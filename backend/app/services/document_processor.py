import os
from typing import Dict, Any
from app.utils.validators import clean_text_content

class DocumentProcessor:
    @staticmethod
    def extract_text_from_file(file_path: str, file_type: str) -> str:
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found: {file_path}")

        extracted_text = ""
        file_type = file_type.lower()

        if file_type == "pdf":
            try:
                import pypdf
                reader = pypdf.PdfReader(file_path)
                pages_text = []
                for idx, page in enumerate(reader.pages):
                    text = page.extract_text() or ""
                    if text.strip():
                        pages_text.append(text)
                extracted_text = "\n\n".join(pages_text)
            except Exception as e:
                # Fallback reading
                with open(file_path, "rb") as f:
                    extracted_text = f.read().decode("utf-8", errors="ignore")

        elif file_type in ["docx", "doc"]:
            try:
                import docx
                doc = docx.Document(file_path)
                paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
                # Also extract tables
                for table in doc.tables:
                    for row in table.rows:
                        row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                        if row_text:
                            paragraphs.append(row_text)
                extracted_text = "\n".join(paragraphs)
            except Exception as e:
                with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                    extracted_text = f.read()

        elif file_type in ["txt", "md", "csv"]:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                extracted_text = f.read()
        else:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                extracted_text = f.read()

        return clean_text_content(extracted_text)

    @staticmethod
    def process_content(text: str, title: str) -> Dict[str, Any]:
        cleaned = clean_text_content(text)
        words = cleaned.split()
        word_count = len(words)
        
        # Simple high-yield summary preview
        summary = cleaned[:400] + ("..." if len(cleaned) > 400 else "")
        return {
            "cleaned_text": cleaned,
            "word_count": word_count,
            "summary": summary
        }

document_processor = DocumentProcessor()
